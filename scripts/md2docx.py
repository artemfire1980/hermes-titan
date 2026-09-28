#!/usr/bin/env python3
"""Конвертер Markdown → DOCX для отчётов исследования"""
import sys
import re
from pathlib import Path

def md_to_docx(md_path: Path, docx_path: Path):
    """Конвертирует Markdown в DOCX используя python-docx"""
    try:
        from docx import Document
        from docx.shared import Pt
    except ImportError:
        print("❌ python-docx не установлен. Запусти: pip3 install python-docx", file=sys.stderr)
        sys.exit(1)

    md_text = md_path.read_text(encoding='utf-8')
    doc = Document()

    # Устанавливаем шрифт по умолчанию
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)

    lines = md_text.split('\n')
    i = 0
    in_frontmatter = False
    in_table = False

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Frontmatter YAML (пропускаем)
        if i == 0 and stripped == '---':
            in_frontmatter = True
            i += 1
            continue
        if in_frontmatter:
            if stripped == '---':
                in_frontmatter = False
            i += 1
            continue

        # Пустая строка
        if not stripped:
            i += 1
            continue

        # Заголовки
        if stripped.startswith('#'):
            match = re.match(r'^(#+)\s+(.+)$', stripped)
            if match:
                level = len(match.group(1))
                text = match.group(2)
                # Ограничение: docx поддерживает heading 1-9
                heading_level = min(level, 9)
                doc.add_heading(text, level=heading_level)
                i += 1
                continue

        # Bullet points
        if stripped.startswith('- ') or stripped.startswith('* ') or stripped.startswith('• '):
            text = stripped[2:].strip()
            # Убираем markdown formatting
            text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
            text = re.sub(r'\*(.+?)\*', r'\1', text)
            text = re.sub(r'`(.+?)`', r'\1', text)
            doc.add_paragraph(text, style='List Bullet')
            i += 1
            continue

        # Таблицы (простой парсинг)
        if stripped.startswith('|') and '|' in stripped[1:]:
            if not in_table:
                in_table = True
                # Создаём таблицу
                table_lines = []
                while i < len(lines) and lines[i].strip().startswith('|'):
                    table_lines.append(lines[i].strip())
                    i += 1

                # Парсим таблицу
                rows_data = []
                for tl in table_lines:
                    cells = [c.strip() for c in tl.split('|')[1:-1]]
                    # Пропускаем separator (---|---|---)
                    if cells and all(re.match(r'^[-:]+$', c) for c in cells if c):
                        continue
                    if cells:
                        rows_data.append(cells)

                if rows_data:
                    ncols = max(len(r) for r in rows_data)
                    table = doc.add_table(rows=len(rows_data), cols=ncols)
                    table.style = 'Light Grid Accent 1'
                    for row_idx, row_data in enumerate(rows_data):
                        for col_idx in range(ncols):
                            cell = table.rows[row_idx].cells[col_idx]
                            cell.text = row_data[col_idx] if col_idx < len(row_data) else ''
                            # Первая строка — жирная
                            if row_idx == 0:
                                for paragraph in cell.paragraphs:
                                    for run in paragraph.runs:
                                        run.bold = True
                in_table = False
                continue
            i += 1
            continue

        # Обычный параграф
        text = stripped
        # Убираем markdown formatting
        text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
        text = re.sub(r'\*(.+?)\*', r'\1', text)
        text = re.sub(r'`(.+?)`', r'\1', text)
        text = re.sub(r'\[([^\]]+)\]\([^\)]+\)', r'\1', text)  # [text](url) -> text
        doc.add_paragraph(text)
        i += 1

    doc.save(str(docx_path))
    print(f"✅ Конвертировано: {md_path.name} → {docx_path.name}")

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: md2docx.py <file.md> [output.docx]")
        sys.exit(1)

    md_path = Path(sys.argv[1])
    if not md_path.exists():
        print(f"❌ Файл не найден: {md_path}")
        sys.exit(1)

    if len(sys.argv) >= 3:
        docx_path = Path(sys.argv[2])
    else:
        docx_path = md_path.with_suffix('.docx')

    md_to_docx(md_path, docx_path)
